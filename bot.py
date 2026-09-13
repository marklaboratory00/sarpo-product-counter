import os
import io
import re
import sqlite3
import logging
import threading
from datetime import datetime
from http.server import BaseHTTPRequestHandler, HTTPServer

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup, InputFile
from telegram.ext import (
    Application, CommandHandler, MessageHandler, CallbackQueryHandler,
    ContextTypes, filters,
)

logging.basicConfig(format="%(asctime)s %(levelname)s %(name)s %(message)s", level=logging.INFO)
log = logging.getLogger("sarpo-bot")

# ---- settings ----
TOKEN = os.environ["BOT_TOKEN"]
ADMIN_ID = int(os.environ["ADMIN_ID"])
DB_PATH = os.environ.get("DB_PATH", "bot.db")          # Railway'da /data/bot.db bo'lsin (Volume!)
IMPORT_PATH = os.path.join(os.path.dirname(DB_PATH) or ".", "import.xlsx")

SIZES = ["XL", "L", "M", "S"]

# Filial -> qaysi brendlarni sotadi
BRANCH_BRANDS = {
    "Yunusobod":      ["Sarpo"],
    "High Town Mall": ["Basic"],
    "Park in Mall":   ["Basic"],
    "General Uzakov": ["Basic", "Sarpo"],
    "Sergeli":        ["Basic", "Sarpo"],
}
BRANCHES = list(BRANCH_BRANDS.keys())

# Import faylida (xlsx) har brend varag'idagi filial ustunlari tartibi
BRAND_BRANCHCOLS = {
    "Basic": ["High Town Mall", "Park in Mall", "General Uzakov", "Sergeli"],
    "Sarpo": ["Yunusobod", "General Uzakov", "Sergeli"],
}
STOCK_START_COL = 9   # I ustun
ART_COL, NAME_COL, IMG_COL = 7, 5, 1


# =========================================================
#  DATABASE
# =========================================================
def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = db()
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            position INTEGER, article TEXT, name TEXT, brand TEXT, file_id TEXT
        );
        CREATE TABLE IF NOT EXISTS counts (
            branch TEXT, product_id INTEGER,
            xl INTEGER, l INTEGER, m INTEGER, s INTEGER,
            reviewed INTEGER DEFAULT 0,
            user_id INTEGER, username TEXT, updated_at TEXT,
            PRIMARY KEY (branch, product_id)
        );
        CREATE TABLE IF NOT EXISTS sessions (
            user_id INTEGER PRIMARY KEY, branch TEXT, idx INTEGER
        );
        """
    )
    conn.commit()
    conn.close()


def brand_brands(branch):
    return BRANCH_BRANDS.get(branch, [])


def products_for_branch(branch):
    brands = brand_brands(branch)
    if not brands:
        return []
    conn = db()
    q = "SELECT * FROM products WHERE brand IN (%s) ORDER BY position" % ",".join("?" * len(brands))
    rows = conn.execute(q, brands).fetchall()
    conn.close()
    return rows


def products_count():
    conn = db()
    n = conn.execute("SELECT COUNT(*) c FROM products").fetchone()["c"]
    conn.close()
    return n


def set_file_id(product_id, file_id):
    conn = db()
    conn.execute("UPDATE products SET file_id=? WHERE id=?", (file_id, product_id))
    conn.commit()
    conn.close()


def get_count(branch, product_id):
    conn = db()
    row = conn.execute("SELECT * FROM counts WHERE branch=? AND product_id=?", (branch, product_id)).fetchone()
    conn.close()
    return row


def save_count(branch, product_id, nums, reviewed, user_id, username):
    conn = db()
    conn.execute(
        """INSERT INTO counts(branch,product_id,xl,l,m,s,reviewed,user_id,username,updated_at)
           VALUES(?,?,?,?,?,?,?,?,?,?)
           ON CONFLICT(branch,product_id) DO UPDATE SET
             xl=excluded.xl,l=excluded.l,m=excluded.m,s=excluded.s,
             reviewed=excluded.reviewed,user_id=excluded.user_id,
             username=excluded.username,updated_at=excluded.updated_at""",
        (branch, product_id, nums[0], nums[1], nums[2], nums[3], reviewed,
         user_id, username, datetime.now().isoformat(timespec="seconds")),
    )
    conn.commit()
    conn.close()


def reviewed_count(branch):
    conn = db()
    n = conn.execute("SELECT COUNT(*) c FROM counts WHERE branch=? AND reviewed=1", (branch,)).fetchone()["c"]
    conn.close()
    return n


def get_session(user_id):
    conn = db()
    row = conn.execute("SELECT * FROM sessions WHERE user_id=?", (user_id,)).fetchone()
    conn.close()
    return row


def set_session(user_id, branch, idx):
    conn = db()
    conn.execute(
        """INSERT INTO sessions(user_id,branch,idx) VALUES(?,?,?)
           ON CONFLICT(user_id) DO UPDATE SET branch=excluded.branch, idx=excluded.idx""",
        (user_id, branch, idx),
    )
    conn.commit()
    conn.close()


def clear_counts():
    conn = db()
    conn.execute("DELETE FROM counts")
    conn.execute("DELETE FROM sessions")
    conn.commit()
    conn.close()


def who_label(user):
    return f"@{user.username}" if user.username else user.full_name


# =========================================================
#  IMPORT (xlsx -> baza)
# =========================================================
def do_import():
    wb = load_workbook(IMPORT_PATH)
    conn = db()
    pos = conn.execute("SELECT COALESCE(MAX(position),0) m FROM products").fetchone()["m"]
    added = updated = counts_added = 0
    for ws in wb.worksheets:
        brand = ws.title.strip()
        if brand not in BRAND_BRANCHCOLS:
            continue
        branches = BRAND_BRANCHCOLS[brand]
        for r in range(2, ws.max_row + 1):
            art = ws.cell(r, ART_COL).value
            if art in (None, ""):
                continue
            art = str(int(art)) if isinstance(art, (int, float)) else str(art).strip()
            name = ws.cell(r, NAME_COL).value or ""
            imgf = ws.cell(r, IMG_COL).value
            url = ""
            if isinstance(imgf, str):
                mm = re.search(r'https://[^"]+', imgf)
                if mm:
                    url = mm.group(0)
            row_ex = conn.execute("SELECT id FROM products WHERE article=? AND brand=?", (art, brand)).fetchone()
            if row_ex:
                pid = row_ex["id"]
                conn.execute("UPDATE products SET name=?, file_id=? WHERE id=?", (name, url, pid))
                updated += 1
            else:
                pos += 1
                cur = conn.execute(
                    "INSERT INTO products(position,article,name,brand,file_id) VALUES(?,?,?,?,?)",
                    (pos, art, name, brand, url),
                )
                pid = cur.lastrowid
                added += 1
            for bi, br in enumerate(branches):
                base = STOCK_START_COL + bi * 4
                vals = [ws.cell(r, base + si).value for si in range(4)]
                if all(v in (None, "") for v in vals):
                    continue
                nums = [int(v) if isinstance(v, (int, float)) else 0 for v in vals]
                conn.execute(
                    """INSERT INTO counts(branch,product_id,xl,l,m,s,reviewed,user_id,username,updated_at)
                       VALUES(?,?,?,?,?,?,0,0,'import',?)
                       ON CONFLICT(branch,product_id) DO UPDATE SET
                         xl=excluded.xl,l=excluded.l,m=excluded.m,s=excluded.s""",
                    (br, pid, nums[0], nums[1], nums[2], nums[3],
                     datetime.now().isoformat(timespec="seconds")),
                )
                counts_added += 1
    conn.commit()
    conn.close()
    return added, updated, counts_added


async def on_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return
    doc = update.message.document
    if not doc or not doc.file_name.lower().endswith((".xlsx", ".xls")):
        await update.message.reply_text("Excel (.xlsx) fayl yuboring.")
        return
    f = await doc.get_file()
    await f.download_to_drive(IMPORT_PATH)
    await update.message.reply_text("📥 Fayl qabul qilindi. Import uchun /import bosing.")


async def cmd_import(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return
    if not os.path.exists(IMPORT_PATH):
        await update.message.reply_text("Avval Excel faylni yuboring, keyin /import.")
        return
    await update.message.reply_text("⏳ Import boshlandi…")
    try:
        added, updated, counts_added = do_import()
        await update.message.reply_text(
            f"✅ Import tayyor.\nYangi mahsulot: {added}\nYangilangan: {updated}\n"
            f"Qoldiq yozuvlari: {counts_added}\n\nEndi xodimlar /start bilan tekshirishi mumkin."
        )
    except Exception as e:
        log.exception("import error")
        await update.message.reply_text(f"❌ Import xatosi: {e}")


# =========================================================
#  ADMIN buyruqlari
# =========================================================
async def on_photo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    # yangi mahsulotni qo'lda qo'shish: rasm + "artikul brend nomi"
    user = update.effective_user
    if user.id != ADMIN_ID:
        await update.message.reply_text("Faqat administrator uchun.")
        return
    cap = (update.message.caption or "").strip()
    if not cap:
        await update.message.reply_text("Izoh: artikul brend nomi. Masalan:\n22400 Sarpo Anor gullari - oq")
        return
    parts = cap.split(None, 2)
    art = parts[0]
    brand = parts[1] if len(parts) > 1 and parts[1] in BRAND_BRANCHCOLS else "Sarpo"
    name = parts[2].strip() if len(parts) > 2 else (parts[1] if len(parts) > 1 else "")
    file_id = update.message.photo[-1].file_id
    conn = db()
    pos = conn.execute("SELECT COALESCE(MAX(position),0) m FROM products").fetchone()["m"] + 1
    conn.execute("INSERT INTO products(position,article,name,brand,file_id) VALUES(?,?,?,?,?)",
                 (pos, art, name, brand, file_id))
    conn.commit()
    conn.close()
    await update.message.reply_text(f"✅ Qo'shildi: {art} · {brand} · {name}")


async def cmd_report(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return
    lines = []
    for b in BRANCHES:
        total = len(products_for_branch(b))
        done = reviewed_count(b)
        lines.append(f"{b}: {done}/{total}")
    await update.message.reply_text("Tekshirilgan holat:\n" + "\n".join(lines))


async def cmd_newcount(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return
    kb = InlineKeyboardMarkup([[InlineKeyboardButton("Ha, yangi hisob", callback_data="newcount")]])
    await update.message.reply_text("Barcha qoldiqlar nollanadi (katalog qoladi). Davom etamizmi?", reply_markup=kb)


async def cmd_export(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Qoldiqlar"
    ws.append(["Filial", "Brend", "Artikul", "Nomi", "XL", "L", "M", "S", "Tekshirilgan", "Sana"])
    for c in ws[1]:
        c.font = Font(bold=True)
    conn = db()
    rows = conn.execute(
        """SELECT c.branch, p.brand, p.article, p.name, c.xl,c.l,c.m,c.s, c.reviewed, c.updated_at
           FROM counts c JOIN products p ON p.id=c.product_id
           ORDER BY c.branch, p.position"""
    ).fetchall()
    conn.close()
    for r in rows:
        ws.append([r["branch"], r["brand"], r["article"], r["name"],
                   r["xl"], r["l"], r["m"], r["s"],
                   "ha" if r["reviewed"] else "yo'q", r["updated_at"]])
    for i, w in enumerate([16, 8, 10, 26, 5, 5, 5, 5, 12, 20], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    await update.message.reply_document(
        document=InputFile(bio, filename=f"qoldiqlar_{datetime.now():%Y-%m-%d}.xlsx"))


# =========================================================
#  XODIM: tekshirish jarayoni
# =========================================================
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if products_count() == 0:
        await update.message.reply_text("Katalog bo'sh. Administrator import qilishini kuting.")
        return
    kb = [[InlineKeyboardButton(b, callback_data=f"b:{i}")] for i, b in enumerate(BRANCHES)]
    await update.message.reply_text("Filialni tanlang:", reply_markup=InlineKeyboardMarkup(kb))


async def send_product(context: ContextTypes.DEFAULT_TYPE, user_id: int):
    s = get_session(user_id)
    if not s:
        return
    branch = s["branch"]
    prods = products_for_branch(branch)
    total = len(prods)
    idx = s["idx"]
    if idx >= total:
        await context.bot.send_message(user_id, f"✅ Tayyor! {branch}: {total} tadan {total} tasi tekshirildi. Rahmat!")
        try:
            await context.bot.send_message(ADMIN_ID, f"✅ {branch} — tekshiruv yakunlandi ({total} ta).")
        except Exception:
            pass
        return
    p = prods[idx]
    ex = get_count(branch, p["id"])
    head = f"{p['article']} · {p['brand']} · {p['name']}\nMahsulot {idx+1} / {total}"
    if ex:
        cur = f"XL:{ex['xl']}  L:{ex['l']}  M:{ex['m']}  S:{ex['s']}"
        caption = f"{head}\n\nHozirgi qoldiq: {cur}\nTo'g'ri bo'lsa — ✅. Aks holda yangi 4 ta son yozing: {' '.join(SIZES)}"
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton("✅ To'g'ri (keyingisi)", callback_data="ok")],
            [InlineKeyboardButton("← Orqaga", callback_data="back")],
        ])
    else:
        caption = f"{head}\n\nQoldiqni yozing: {' '.join(SIZES)}\nMasalan: 2 3 1 0"
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton("∅ Qoldiq yo'q", callback_data="zero")],
            [InlineKeyboardButton("← Orqaga", callback_data="back")],
        ])
    photo = p["file_id"]
    try:
        msg = await context.bot.send_photo(user_id, photo=photo, caption=caption, reply_markup=kb)
        # URL bo'lsa -> Telegram file_id ni saqlab qo'yamiz (keyin Drive kerak bo'lmaydi)
        if isinstance(photo, str) and photo.startswith("http") and msg.photo:
            set_file_id(p["id"], msg.photo[-1].file_id)
    except Exception as e:
        log.warning("send_photo failed (%s): %s", p["article"], e)
        await context.bot.send_message(user_id, caption + "\n\n(rasm yuklanmadi)", reply_markup=kb)


async def on_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    s = get_session(user.id)
    if not s:
        return
    parts = update.message.text.replace(",", " ").split()
    if len(parts) != 4 or not all(x.isdigit() for x in parts):
        await update.message.reply_text(f"4 ta son kerak: {' '.join(SIZES)}. Masalan: 2 3 1 0")
        return
    nums = [int(x) for x in parts]
    context.user_data["pending"] = nums
    summary = " · ".join(f"{sz}:{n}" for sz, n in zip(SIZES, nums))
    kb = InlineKeyboardMarkup([[
        InlineKeyboardButton("✅ Saqlash", callback_data="save"),
        InlineKeyboardButton("✏️ Qaytadan", callback_data="redo"),
    ]])
    await update.message.reply_text(summary, reply_markup=kb)


async def advance(context, user, nums=None, accept=False):
    s = get_session(user.id)
    if not s:
        return
    branch = s["branch"]
    prods = products_for_branch(branch)
    idx = s["idx"]
    if idx >= len(prods):
        return
    p = prods[idx]
    if accept:
        ex = get_count(branch, p["id"])
        nums = [ex["xl"], ex["l"], ex["m"], ex["s"]] if ex else [0, 0, 0, 0]
    save_count(branch, p["id"], nums, 1, user.id, user.username or user.full_name)
    set_session(user.id, branch, idx + 1)
    await send_product(context, user.id)


async def on_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    data = query.data
    user = query.from_user
    await query.answer()

    if data.startswith("b:"):
        branch = BRANCHES[int(data.split(":")[1])]
        total = len(products_for_branch(branch))
        done = reviewed_count(branch)
        if total and done >= total:
            kb = InlineKeyboardMarkup([[InlineKeyboardButton("🔄 Qaytadan", callback_data=f"restart:{BRANCHES.index(branch)}")]])
            await query.edit_message_text(f"{branch}: {total}/{total} tekshirilgan.", reply_markup=kb)
            return
        if 0 < done < total:
            i = BRANCHES.index(branch)
            kb = InlineKeyboardMarkup([
                [InlineKeyboardButton(f"▶️ Davom ({done+1}-dan)", callback_data=f"cont:{i}")],
                [InlineKeyboardButton("🔄 Boshidan", callback_data=f"restart:{i}")],
            ])
            await query.edit_message_text(f"{branch}: {done}/{total} tekshirilgan.", reply_markup=kb)
            return
        set_session(user.id, branch, 0)
        await query.edit_message_text(f"Filial: {branch}")
        await send_product(context, user.id)
        return

    if data.startswith("cont:"):
        branch = BRANCHES[int(data.split(":")[1])]
        set_session(user.id, branch, reviewed_count(branch))
        await query.edit_message_text(f"Davom: {branch}")
        await send_product(context, user.id)
        return
    if data.startswith("restart:"):
        branch = BRANCHES[int(data.split(":")[1])]
        set_session(user.id, branch, 0)
        await query.edit_message_text(f"Boshidan: {branch}")
        await send_product(context, user.id)
        return

    try:
        await query.edit_message_reply_markup(reply_markup=None)
    except Exception:
        pass

    if data == "ok":
        await advance(context, user, accept=True)
    elif data == "zero":
        await advance(context, user, nums=[0, 0, 0, 0])
    elif data == "save":
        nums = context.user_data.pop("pending", None)
        if nums is None:
            await context.bot.send_message(user.id, "Avval 4 ta son kiriting.")
            return
        await advance(context, user, nums=nums)
    elif data == "redo":
        context.user_data.pop("pending", None)
        await context.bot.send_message(user.id, "Qaytadan 4 ta son yozing: " + " ".join(SIZES))
    elif data == "back":
        s = get_session(user.id)
        if s:
            set_session(user.id, s["branch"], max(0, s["idx"] - 1))
            await send_product(context, user.id)
    elif data == "newcount":
        if user.id == ADMIN_ID:
            clear_counts()
            await context.bot.send_message(user.id, "Yangi hisob boshlandi. Qoldiqlar nollandi.")


# =========================================================
#  Health server + startup
# =========================================================
class Health(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200); self.end_headers(); self.wfile.write(b"ok")
    def log_message(self, *a): pass


def run_health():
    HTTPServer(("0.0.0.0", int(os.environ.get("PORT", "8080"))), Health).serve_forever()


def main():
    init_db()
    threading.Thread(target=run_health, daemon=True).start()
    app = Application.builder().token(TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("import", cmd_import))
    app.add_handler(CommandHandler("report", cmd_report))
    app.add_handler(CommandHandler("export", cmd_export))
    app.add_handler(CommandHandler("newcount", cmd_newcount))
    app.add_handler(CallbackQueryHandler(on_callback))
    app.add_handler(MessageHandler(filters.PHOTO, on_photo))
    app.add_handler(MessageHandler(filters.Document.ALL, on_document))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, on_text))
    log.info("Bot started")
    app.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
